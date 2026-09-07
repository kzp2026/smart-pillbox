import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts" / "09_evaluate_design_scheme.py"
sys.path.insert(0, str(SCRIPT_PATH.parent))
SPEC = importlib.util.spec_from_file_location("evaluate_design_scheme", SCRIPT_PATH)
evaluation = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(evaluation)


class EvaluationHonestyTests(unittest.TestCase):
    def test_empty_materials_have_zero_completeness_score(self):
        self.assertEqual(evaluation.calculate_score("需求匹配度", pd.DataFrame(), ""), 0)

    def test_evaluation_rows_identify_rule_self_check_source(self):
        table = evaluation.build_evaluation_table("产品", pd.DataFrame(), "")

        self.assertIn("来源", table.columns)
        self.assertEqual(set(table["来源"]), {"自动规则自检（非专家评分）"})
        self.assertEqual(set(table["分值"]), {0})

    def test_docx_and_prompt_disclose_no_expert_or_feasibility_claim(self):
        table = evaluation.build_evaluation_table("产品", pd.DataFrame(), "")
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "summary.docx"
            evaluation.save_summary_docx("产品", table, output_path)

            from docx import Document

            document_text = "\n".join(paragraph.text for paragraph in Document(output_path).paragraphs)

        prompt = evaluation.build_optimization_prompt("产品", table, pd.DataFrame())
        for text in (document_text, prompt):
            self.assertIn("自动规则自检（非专家评分）", text)
            self.assertIn("真实专家/用户实验", text)
            self.assertNotIn("工程可行性：系统", text)

    def test_empty_material_report_does_not_claim_outputs_were_generated(self):
        table = evaluation.build_evaluation_table("产品", pd.DataFrame(), "")
        with tempfile.TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "summary.docx"
            evaluation.save_summary_docx("产品", table, output_path)

            from docx import Document

            document_text = "\n".join(paragraph.text for paragraph in Document(output_path).paragraphs)

        self.assertIn("流程预期材料（实际产物以归档为准）", document_text)
        self.assertNotIn("已生成材料", document_text)


if __name__ == "__main__":
    unittest.main()
