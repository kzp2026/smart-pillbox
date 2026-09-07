import json
import hashlib
import subprocess
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from experiment.evaluation.materials import prepare_review_materials
from experiment.evaluation.templates import import_annotations, import_mapping_reviews, read_review_file


class ReviewMaterialsTests(unittest.TestCase):
    def _run(self, root: Path, *, generation: bool = False) -> Path:
        run = root / "run"
        clean = run / "stages" / "clean" / "c1"
        req = run / "stages" / "requirements" / "r1"
        mapping = run / "stages" / "mapping" / "m1"
        clean.mkdir(parents=True)
        req.mkdir(parents=True)
        mapping.mkdir(parents=True)
        comments = [{
            "comment_id": "C1", "original_comment": "给老人用，字再大些", "cleaned_comment": "给老人用,字再大些",
            "date": "2026-03-01", "product_version": "升级款", "source_channel": "京东（待核对）",
        }]
        requirements = [{"requirement_id": "R1", "requirement_name": "清晰阅读", "requirement_description": "字号可辨识"}]
        mappings = [{
            "mapping_id": "M1", "requirement_id": "R1", "requirement_name": "清晰阅读",
            "function_id": "F1", "function_name": "显示信息", "structure_id": "S1", "structure_name": "大字屏",
            "source_comment_ids": ["C1"], "evidence_spans": [{"comment_id": "C1", "text": "字再大些"}],
            "mapping_reason": {"requirement_to_function": "候选理由1", "function_to_structure": "候选理由2"},
            "derivation_method": "rules", "review_status": "pending_review", "reviewer_id": "", "reviewer_note": "",
        }]
        (clean / "comments.json").write_text(json.dumps(comments, ensure_ascii=False), encoding="utf-8")
        (req / "requirements.json").write_text(json.dumps(requirements, ensure_ascii=False), encoding="utf-8")
        (mapping / "mappings.json").write_text(json.dumps(mappings, ensure_ascii=False), encoding="utf-8")
        artifacts = ["stages/clean/c1/comments.json", "stages/requirements/r1/requirements.json", "stages/mapping/m1/mappings.json"]
        if generation:
            gen = run / "stages" / "generation" / "g1" / "blind"
            gen.mkdir(parents=True)
            (gen / "blind_schemes.json").write_text(json.dumps([{"scheme_id": "A", "content": "方案"}], ensure_ascii=False), encoding="utf-8")
            artifacts.append("stages/generation/g1/blind/blind_schemes.json")
        stages = {}
        for stage, rel in (("clean", artifacts[0]), ("requirements", artifacts[1]), ("mapping", artifacts[2])):
            stages[stage] = {"status": "completed", "outputs": {rel: hashlib.sha256((run / rel).read_bytes()).hexdigest()}}
        if generation:
            rel = artifacts[-1]; stages["generation"] = {"status": "completed", "outputs": {rel: hashlib.sha256((run / rel).read_bytes()).hexdigest()}}
        (run / "run_manifest.json").write_text(json.dumps({"run_id": "RUN1", "artifacts": artifacts, "stages": stages}), encoding="utf-8")
        return run

    def test_exports_pre_generation_materials_without_leaking_system_labels(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); run = self._run(root)
            result = prepare_review_materials(run, root / "materials")
            self.assertEqual("ready_with_missing_items", result["status"])
            self.assertIn("generation_schemes", result["missing_items"])
            initial = load_workbook(root / "materials" / "01_initial_free_annotation.xlsx")
            headers = [c.value for c in initial.active[1]]
            self.assertEqual(["comment_id", "评论原文", "使用问题", "自由需求描述", "无法判断", "annotator_id", "annotator_role", "annotated_at", "notes"], headers)
            flat = " ".join(str(v) for row in initial.active.values for v in row if v is not None)
            self.assertNotIn("R1", flat)
            self.assertNotIn("清晰阅读", flat)
            source = load_workbook(root / "materials" / "00_data_source_check.xlsx")
            self.assertEqual(["平台", "产品标识", "采集时间", "评论时间范围", "采集方式", "筛选规则", "来源凭据", "核对人", "核对时间", "备注"], [c.value for c in source.active[1]])
            elderly = load_workbook(root / "materials" / "02_elderly_context_candidates.xlsx")
            self.assertIn("候选提示（非人工结论）", [c.value for c in elderly.active[1]])
            self.assertEqual("明确文本支持", elderly.active.cell(2, 3).value)
            disclosure = (root / "materials" / "BOUNDARY_DISCLOSURE.md").read_text(encoding="utf-8")
            self.assertIn("rules 已查看全部数据", disclosure)
            self.assertIn("不能声称独立验证", disclosure)

    def test_mapping_and_review_templates_are_operational_and_compatible(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); run = self._run(root, generation=True)
            prepare_review_materials(run, root / "materials")
            mapping = load_workbook(root / "materials" / "04_mapping_review.xlsx")
            headers = [c.value for c in mapping.active[1]]
            for expected in ("mapping_id", "requirement_id", "function_id", "structure_id", "evidence_spans", "mapping_reason", "derivation_method", "knowledge_source", "review_decision", "review_status", "reviewer_id", "reviewer_note", "reviewed_at"):
                self.assertIn(expected, headers)
            row = 2
            mapping.active.cell(row, headers.index("review_decision") + 1, "agree")
            mapping.active.cell(row, headers.index("reviewer_id") + 1, "REV-01")
            mapping.active.cell(row, headers.index("reviewer_note") + 1, "证据与关系理由一致")
            mapping.active.cell(row, headers.index("reviewed_at") + 1, "2026-09-07T10:00:00+08:00")
            mapping.save(root / "materials" / "04_mapping_review.xlsx")
            normalized = import_mapping_reviews(root / "materials" / "04_mapping_review.xlsx")
            self.assertEqual("approved", normalized[0]["review_status"])
            rubric = load_workbook(root / "materials" / "05_blind_review.xlsx")["评分说明"]
            self.assertEqual(7 * 5 + 1, rubric.max_row)
            changes = [c.value for c in load_workbook(root / "materials" / "06_v1_v2_changes.xlsx").active[1]]
            self.assertIn("old_value", changes); self.assertIn("new_value", changes)
            self.assertIn("linked_review_comment", changes); self.assertIn("evaluated_at", changes)
            blind_path = root / "materials" / "05_blind_review.xlsx"
            blind = load_workbook(blind_path); headers = [c.value for c in blind["填写表"][1]]
            values = {"匿名评委 ID":"REV-01","评委背景":"工业设计","匿名方案 ID":"A","方案内容":"方案","优点":"反馈明确","问题":"按钮小","建议":"增大按钮","评价时间":"2026-09-07T10:00:00+08:00"}
            values.update({dimension: 3 for dimension in __import__('experiment.evaluation.templates', fromlist=['DIMENSIONS']).DIMENSIONS})
            for column, header in enumerate(headers, 1): blind["填写表"].cell(2, column, values.get(header))
            blind.save(blind_path)
            self.assertEqual("REV-01", read_review_file(blind_path)[0]["reviewer_id"])

    def test_initial_import_supports_free_form_and_legacy_coded_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); run = self._run(root)
            prepare_review_materials(run, root / "materials")
            path = root / "materials" / "01_initial_free_annotation.xlsx"
            book = load_workbook(path); sheet = book.active
            sheet.cell(2, 3, "字号太小"); sheet.cell(2, 4, "希望字体更大"); sheet.cell(2, 5, "否")
            sheet.cell(2, 6, "ANN-01"); sheet.cell(2, 7, "研究助理"); sheet.cell(2, 8, "2026-09-06T10:00:00+08:00")
            book.save(path)
            rows = import_annotations(path, {"C1"}, {"R1"})
            self.assertEqual("希望字体更大", rows[0]["free_requirement"])
            self.assertNotIn("requirement_id", rows[0])

    def test_cli_prints_missing_items_and_import_commands(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); run = self._run(root); out = root / "out"
            proc = subprocess.run([
                str(Path(".venv-research/Scripts/python.exe")), "scripts/prepare_review.py",
                "--run", str(run), "--output", str(out),
            ], cwd=Path(__file__).resolve().parents[1], text=True, capture_output=True, encoding="utf-8")
            self.assertEqual(0, proc.returncode, proc.stderr)
            self.assertIn("材料缺失清单", proc.stdout)
            self.assertIn("--annotations", proc.stdout)
            self.assertIn("--mapping-review", proc.stdout)

    def test_uses_only_manifest_current_stage_and_rejects_hash_mismatch(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); run = self._run(root)
            stale = run / "stages" / "clean" / "stale"; stale.mkdir(parents=True)
            (stale / "comments.json").write_text('[{"comment_id":"STALE","original_comment":"错误"}]', encoding="utf-8")
            prepare_review_materials(run, root / "first")
            package = json.loads((root / "first" / "materials_manifest.json").read_text(encoding="utf-8"))
            self.assertIn("sha256", package["source_artifacts"]["clean"])
            self.assertIn("01_initial_free_annotation.xlsx", package["exported_file_sha256"])
            with self.assertRaisesRegex(ValueError, "must be empty"):
                prepare_review_materials(run, root / "first")
            values = list(load_workbook(root / "first" / "01_initial_free_annotation.xlsx").active.values)
            self.assertEqual("C1", values[1][0])
            current = run / "stages" / "clean" / "c1" / "comments.json"
            current.write_text("[]", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "哈希不一致"):
                prepare_review_materials(run, root / "tampered")


if __name__ == "__main__":
    unittest.main()
