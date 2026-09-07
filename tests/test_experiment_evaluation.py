import math
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from experiment.evaluation.loop import validate_changes
from experiment.evaluation.statistics import (
    DIMENSIONS,
    compare_versions,
    summarize_reviews,
    validate_reviews,
)
from experiment.evaluation.templates import create_templates, import_annotations, read_review_file


def review(reviewer, scheme, scores, *, when="2026-09-05T12:00:00+08:00"):
    row = {
        "reviewer_id": reviewer,
        "background": "工业设计",
        "scheme_id": scheme,
        "evaluated_at": when,
        "strengths": "层级清晰",
        "issues": "按钮偏小",
        "suggestions": "增大按钮",
    }
    row.update(dict(zip(DIMENSIONS, scores)))
    return row


class TemplateTests(unittest.TestCase):
    def test_creates_six_anonymous_workbooks_and_preserves_mapping_fields(self):
        comments = [{"comment_id": "C1", "cleaned_comment": "希望字大"}]
        requirements = [{"requirement_id": "R1", "name": "大字"}]
        mappings = [{"requirement_id": "R1", "function_id": "F1", "structure_id": "S1", "confidence": 0.8}]
        schemes = [{"scheme_id": "S-A", "content": "匿名内容", "version": "V1"}]
        with tempfile.TemporaryDirectory() as tmp:
            paths = create_templates(Path(tmp), comments, requirements, mappings, schemes)
            self.assertEqual(6, len(paths))
            self.assertTrue(all(path.suffix == ".xlsx" and path.exists() for path in paths))
            mapping_book = load_workbook(next(p for p in paths if "expert_mapping" in p.name))
            rows = list(mapping_book.active.values)
            self.assertIn("confidence", rows[0])
            self.assertEqual("pending_review", rows[1][rows[0].index("review_status")])
            blind_book = load_workbook(next(p for p in paths if "abc_blind_review" in p.name))
            blind_values = list(blind_book.active.values)
            self.assertIn("匿名评委 ID", blind_values[0])
            self.assertTrue(set(DIMENSIONS).issubset(blind_values[0]))
            flat = " ".join("" if v is None else str(v) for row in blind_values for v in row)
            self.assertNotIn("V1", flat)
            self.assertNotIn("group", flat.lower())

    def test_serializes_nested_mapping_and_neutralizes_formula_cells(self):
        mappings = [{"requirement_id": "R1", "function_id": "F1", "structure_id": "S1", "evidence": {"source": "=CMD()"}, "tags": ["a", "b"]}]
        with tempfile.TemporaryDirectory() as tmp:
            path = next(p for p in create_templates(Path(tmp), [], [], mappings, []) if "expert_mapping" in p.name)
            rows = list(load_workbook(path, data_only=False).active.values)
            evidence = rows[1][rows[0].index("evidence")]
            self.assertTrue(evidence.startswith("{"))
            self.assertIn("'=CMD()", evidence)

    def test_reads_chinese_review_sheet_and_skips_fully_empty_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = next(p for p in create_templates(Path(tmp), [], [], [], [{"scheme_id": "SCH-A", "content": "内容", "version": "V1"}, {"scheme_id": "SCH-B", "content": "未评价", "version": "V1"}]) if "abc_blind" in p.name)
            book = load_workbook(path)
            sheet = book.active
            values = {"匿名评委 ID": "REV-01", "评委背景": "设计", "匿名方案 ID": "SCH-A", "方案内容": "内容", "优点": "清晰", "问题": "按钮小", "建议": "放大", "评价时间": "2026-09-05T12:00:00+08:00"}
            values.update({dimension: 3 for dimension in DIMENSIONS})
            for column, header in enumerate([cell.value for cell in sheet[1]], 1):
                sheet.cell(2, column, values.get(header))
            sheet.append([None] * sheet.max_column)
            book.save(path)
            rows = read_review_file(path)
            self.assertEqual(1, len(rows))
            self.assertEqual("REV-01", rows[0]["reviewer_id"])
            self.assertIsInstance(rows[0][DIMENSIONS[0]], int)
            sheet.cell(3, 1, "REV-02")
            book.save(path)
            with self.assertRaises(ValueError):
                read_review_file(path)

    def test_imports_multi_annotator_rows_and_validates_ids_and_no_requirement(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = next(p for p in create_templates(Path(tmp), [{"comment_id": "C1", "cleaned_comment": "文字小"}, {"comment_id": "C2", "cleaned_comment": "未标注"}], [], [], []) if "requirement_annotation" in p.name)
            book = load_workbook(path); sheet = book.active
            sheet.append(["C1", "文字小", "R1", "ANN-02", "2026-09-05T12:00:00+08:00", ""])
            sheet.append(["C1", "文字小", "NO_REQUIREMENT", "ANN-03", "2026-09-05T12:01:00+08:00", "无明确需求"])
            sheet.cell(2, 3, "R1"); sheet.cell(2, 4, "ANN-01"); sheet.cell(2, 5, "2026-09-05T12:00:00+08:00")
            book.save(path)
            rows = import_annotations(path, {"C1", "C2"}, {"R1"})
            self.assertEqual(3, len(rows))
            self.assertIsNone(rows[2]["requirement_id"])
            sheet.cell(2, 1, "UNKNOWN")
            book.save(path)
            with self.assertRaises(ValueError):
                import_annotations(path, {"C1"}, {"R1"})


class ValidationTests(unittest.TestCase):
    def test_accepts_complete_rows_and_normalizes_scores(self):
        result = validate_reviews([review("REV-01", "SCH-A", [1, 2, 3, 4, 5, 4, 3])], {"SCH-A"})
        self.assertEqual(5, result[0]["结构合理性"])

    def test_rejects_invalid_scores_ids_times_duplicates_and_unknown_schemes(self):
        cases = []
        bad = review("REV-01", "SCH-A", [1, 2, 3, 4, 5, 4, 3]); bad[DIMENSIONS[0]] = 2.5; cases.append(([bad], {"SCH-A"}))
        bad = review("human name", "SCH-A", [1] * 7); cases.append(([bad], {"SCH-A"}))
        bad = review("REV-01", "SCH-A", [1] * 7, when="yesterday"); cases.append(([bad], {"SCH-A"}))
        cases.append(([review("REV-01", "SCH-X", [1] * 7)], {"SCH-A"}))
        duplicate = review("REV-01", "SCH-A", [1] * 7); cases.append(([duplicate, dict(duplicate)], {"SCH-A"}))
        bad = review("REV-01", "SCH-A", [1] * 7); bad[DIMENSIONS[0]] = math.nan; cases.append(([bad], {"SCH-A"}))
        for rows, ids in cases:
            with self.subTest(rows=rows):
                with self.assertRaises(ValueError):
                    validate_reviews(rows, ids)


class StatisticsTests(unittest.TestCase):
    def test_summary_matches_hand_calculated_sample_statistics(self):
        rows = [review(f"REV-{i:02d}", "SCH-A", [score] * 7) for i, score in enumerate([1, 2, 3, 4, 5], 1)]
        result = summarize_reviews(rows)
        metric = result["schemes"]["SCH-A"]["dimensions"][DIMENSIONS[0]]
        self.assertEqual(5, metric["n"])
        self.assertAlmostEqual(3.0, metric["mean"])
        self.assertAlmostEqual(math.sqrt(2.5), metric["sample_std"])
        self.assertEqual(3, metric["median"])
        self.assertAlmostEqual(1.0367568, metric["ci95"][0], places=5)
        self.assertAlmostEqual(4.9632432, metric["ci95"][1], places=5)
        self.assertFalse(result["simulated"])

    def test_summary_reports_icc_only_for_complete_crossed_design(self):
        complete = [
            review("REV-01", "SCH-A", [1] * 7), review("REV-01", "SCH-B", [2] * 7),
            review("REV-02", "SCH-A", [3] * 7), review("REV-02", "SCH-B", [4] * 7),
        ]
        self.assertEqual("ok", summarize_reviews(complete)["icc_2_1"]["status"])
        self.assertEqual("insufficient_complete_crossed_design", summarize_reviews(complete[:-1])["icc_2_1"]["status"])

    def test_compare_versions_uses_pair_aggregate_and_hand_calculated_effect(self):
        rows = []
        pairs = []
        for pair_no in range(1, 6):
            v1, v2 = f"A{pair_no}", f"B{pair_no}"
            pairs.append({"pair_id": f"P{pair_no}", "v1_scheme_id": v1, "v2_scheme_id": v2})
            for reviewer in ("REV-01", "REV-02"):
                rows += [review(reviewer, v1, [2] * 7), review(reviewer, v2, [3] * 7)]
        result = compare_versions(rows, pairs)
        metric = result["dimensions"][DIMENSIONS[0]]
        self.assertEqual(5, metric["n_pairs"])
        self.assertEqual(1.0, metric["mean_difference"])
        self.assertEqual(1.0, metric["matched_rank_biserial"])
        self.assertAlmostEqual(0.0625, metric["p_value"])
        self.assertAlmostEqual(0.4375, metric["holm_adjusted_p"])
        self.assertFalse(metric["significant"])

    def test_compare_versions_suppresses_conclusions_for_small_or_simulated_samples(self):
        rows = [review("REV-01", "A", [2] * 7), review("REV-01", "B", [3] * 7)]
        pairs = [{"pair_id": "P1", "v1_scheme_id": "A", "v2_scheme_id": "B"}]
        self.assertFalse(compare_versions(rows, pairs)["can_claim_significance"])
        simulated = compare_versions(rows * 0 + rows, pairs, simulated=True)
        self.assertFalse(simulated["can_claim_real_improvement"])

    def test_compare_versions_rejects_invalid_pair_design_and_counts_only_observed_pairs(self):
        rows = [review("REV-01", "SCH-A", [2] * 7), review("REV-01", "SCH-B", [3] * 7)]
        bad_designs = [
            [{"pair_id": "P1", "v1_scheme_id": "SCH-A", "v2_scheme_id": "SCH-A"}],
            [{"pair_id": "P1", "v1_scheme_id": "SCH-A", "v2_scheme_id": "SCH-B"}, {"pair_id": "P1", "v1_scheme_id": "SCH-C", "v2_scheme_id": "SCH-D"}],
            [{"pair_id": "P1", "v1_scheme_id": "SCH-A", "v2_scheme_id": "SCH-B"}, {"pair_id": "P2", "v1_scheme_id": "SCH-A", "v2_scheme_id": "SCH-C"}],
        ]
        for pairs in bad_designs:
            with self.subTest(pairs=pairs), self.assertRaises(ValueError):
                compare_versions(rows, pairs)
        pairs = [{"pair_id": f"P{i}", "v1_scheme_id": f"A{i}", "v2_scheme_id": f"B{i}"} for i in range(1, 6)]
        result = compare_versions(rows, pairs)
        self.assertEqual(0, result["n_independent_pairs"])
        self.assertFalse(result["can_claim_significance"])


class ChangeLoopTests(unittest.TestCase):
    def test_validates_traceable_change(self):
        reviews = [review("REV-01", "SCH-V1", [3, 4, 4, 4, 4, 4, 4])]
        change = {"change_id": "CH-1", "v1_scheme_id": "SCH-V1", "v2_scheme_id": "SCH-V2", "reviewer_id": "REV-01", "low_dimension": DIMENSIONS[0], "expert_issue": "按钮偏小", "requirement_id": "R1", "function_id": "F1", "structure_id": "ST1", "prompt_field": "controls", "old_value": "small", "new_value": "large", "reason": "提高匹配度"}
        result = validate_changes([change], reviews, [{"requirement_id": "R1"}], [{"requirement_id": "R1", "function_id": "F1", "structure_id": "ST1"}], {"SCH-V1", "SCH-V2"})
        self.assertEqual("CH-1", result[0]["change_id"])

    def test_rejects_untraceable_or_unchanged_change(self):
        reviews = [review("REV-01", "SCH-V1", [4] * 7)]
        change = {"change_id": "CH-1", "v1_scheme_id": "SCH-V1", "v2_scheme_id": "SCH-V2", "reviewer_id": "REV-01", "low_dimension": DIMENSIONS[0], "expert_issue": "不同意见", "requirement_id": "R1", "function_id": "F1", "structure_id": "ST1", "prompt_field": "controls", "old_value": "same", "new_value": "same", "reason": "test"}
        with self.assertRaises(ValueError):
            validate_changes([change], reviews, [{"requirement_id": "R1"}], [{"function_id": "F1", "structure_id": "ST1"}], {"SCH-V1", "SCH-V2"})

    def test_rejects_crossed_mapping_same_generation_and_unknown_prompt_field(self):
        reviews = [review("REV-01", "SCH-V1", [3] * 7)]
        base = {"change_id": "CH-1", "v1_scheme_id": "SCH-V1", "v2_scheme_id": "SCH-V2", "reviewer_id": "REV-01", "low_dimension": DIMENSIONS[0], "expert_issue": "按钮偏小", "requirement_id": "R1", "function_id": "F1", "structure_id": "ST1", "prompt_field": "controls", "old_value": "small", "new_value": "large", "reason": "修复"}
        mappings = [{"requirement_id": "R2", "function_id": "F1", "structure_id": "ST1"}, {"requirement_id": "R1", "function_id": "F2", "structure_id": "ST1"}]
        for update in ({}, {"v2_scheme_id": "SCH-V1"}, {"prompt_field": "__private"}):
            change = dict(base, **update)
            with self.subTest(update=update), self.assertRaises(ValueError):
                validate_changes([change], reviews, [{"requirement_id": "R1"}], mappings, {"SCH-V1", "SCH-V2"})


if __name__ == "__main__":
    unittest.main()
