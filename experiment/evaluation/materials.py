from __future__ import annotations

import json
import hashlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from experiment.evaluation.templates import DIMENSIONS, _safe_scalar
from experiment.pipeline.service import artifact, read_manifest


def _read(path: Path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def _artifact(run: Path, manifest: dict, stage: str, filename: str) -> Path | None:
    """Resolve only the current completed stage output and verify its manifest hash."""
    if manifest.get("stages", {}).get(stage, {}).get("status") != "completed":
        return None
    return artifact(run, stage, filename)


def _book(path: Path, headers: list[str], rows: list[dict], *, instructions: str = "") -> Path:
    book = Workbook(); sheet = book.active; sheet.title = "填写表"
    sheet.append(headers)
    for row in rows:
        sheet.append([_safe_scalar(row.get(header, "")) for header in headers])
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(60, max(14, max(len(str(c.value or "")) for c in column) + 2))
        for cell in column[1:]: cell.alignment = Alignment(vertical="top", wrap_text=True)
    if instructions:
        info = book.create_sheet("填写说明"); info["A1"] = instructions; info["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        info.column_dimensions["A"].width = 110; info.row_dimensions[1].height = 180
    path.parent.mkdir(parents=True, exist_ok=True); book.save(path); return path


def _elderly_candidate(text: str) -> tuple[str, str]:
    explicit = ("老人", "老年", "长辈", "父母", "耳背", "爷爷", "奶奶")
    general = ("字大", "声音大", "简单", "易用", "看清", "听不见")
    if any(word in text for word in explicit): return "明确文本支持", "原文出现老人、长辈或相关身份/能力表述"
    if any(word in text for word in general): return "一般", "原文出现可能相关的可用性表述，但未明确老年场景"
    return "未知", "原文无足够文本支持"


def _review_book(path: Path, schemes: list[dict], title: str) -> Path:
    headers = ["匿名评委 ID", "评委背景", "匿名方案 ID", "方案内容", *DIMENSIONS, "优点", "问题", "建议", "评价时间"]
    _book(path, headers, [{"匿名方案 ID": r.get("scheme_id", ""), "方案内容": r.get("content", "")} for r in schemes], instructions="独立填写，不讨论版本与方法身份。七个维度均须按“评分说明”页的1–5锚点评分；文字意见须可操作。")
    book = __import__("openpyxl").load_workbook(path); rubric = book.create_sheet("评分说明")
    rubric.append(["维度", "分数", "可操作评分说明"])
    anchors = {
        "需求匹配度": ["核心需求无对应设计", "仅覆盖少量需求且对应关系含糊", "覆盖主要需求但有可指出的缺口", "主要需求均有明确对应设计", "逐项需求均有可定位的设计响应与一致证据"],
        "适老化友好性": ["关键文字、提示或操作对老年用户构成阻断", "多项关键可达性问题未处理", "基本可完成但字号、提示或步骤仍需明确修改", "关键视觉、听觉与操作环节清楚且有容错", "关键环节均提供可观察的多通道提示、清晰操作与恢复路径"],
        "功能完整性": ["核心任务无法闭环", "缺少多项完成核心任务所需功能", "核心流程可描述但仍缺一个重要环节", "核心流程功能齐备且输入输出明确", "核心、异常与恢复流程均有明确功能响应"],
        "操作便利性": ["关键任务无法按方案说明完成", "步骤冗长或反馈缺失导致高出错风险", "任务可完成但步骤或反馈有明确改进点", "关键任务步骤精简且反馈清楚", "关键任务步骤、状态反馈、撤销与纠错均清楚可查"],
        "结构合理性": ["部件关系矛盾或无法装配", "布局、连接或维护存在多项明显冲突", "总体关系成立但有明确干涉或维护疑点", "部件布局、连接与维护路径描述一致", "关键部件尺寸关系、连接、拆装与维护路径均有明确说明"],
        "工程可行性": ["存在已知物理冲突或不可实现描述", "材料、工艺、供电或成本中多项无交代", "给出基本实现路径但关键假设待验证", "材料工艺与关键约束有一致说明，仍需样机验证", "材料、工艺、供电、安全和成本假设均明确列出并标注验证方法；此分数不等于工程验证通过"],
        "创新性": ["未识别出相对常见方案的差异", "差异仅为外观或措辞变化且价值不清", "有一个可辨识差异但使用价值待论证", "差异机制与用户问题有明确对应", "多个可辨识机制均说明相对基线的差异、使用价值与待验证假设"],
    }
    for dimension in DIMENSIONS:
        for score, description in enumerate(anchors[dimension], 1): rubric.append([dimension, score, description])
    for cell in rubric[1]: cell.font = Font(bold=True)
    book.save(path); return path


def prepare_review_materials(run_dir: Path, output_dir: Path) -> dict:
    run, output = Path(run_dir).resolve(), Path(output_dir).resolve()
    manifest_path = run / "run_manifest.json"
    if not manifest_path.is_file(): raise ValueError("run directory lacks run_manifest.json")
    if output.exists() and any(output.iterdir()):
        raise ValueError("output directory must be empty; existing review materials will not be overwritten")
    manifest = read_manifest(run)
    comment_path = _artifact(run, manifest, "clean", "comments.json")
    if not comment_path: raise ValueError("clean comments artifact is required")
    comments = _read(comment_path)
    req_path = _artifact(run, manifest, "requirements", "requirements.json")
    map_path = _artifact(run, manifest, "mapping", "mappings.json")
    gen_path = _artifact(run, manifest, "generation", "blind_schemes.json")
    requirements, mappings, schemes = _read(req_path) if req_path else [], _read(map_path) if map_path else [], _read(gen_path) if gen_path else []
    output.mkdir(parents=True, exist_ok=True); files = []
    dates = sorted(str(row.get("date") or "") for row in comments if row.get("date"))
    source_rows = [{"平台": "", "产品标识": "", "采集时间": "", "评论时间范围": f"{dates[0]} 至 {dates[-1]}" if dates else "", "采集方式": "", "筛选规则": "", "来源凭据": "", "核对人": "", "核对时间": "", "备注": "未知项留空，须由核对人员补录"}]
    files.append(_book(output / "00_data_source_check.xlsx", list(source_rows[0]), source_rows, instructions="逐项核对数据来源。未知信息必须留空，不得根据文件名或系统提示推断。来源凭据填写可复核的订单、页面、导出记录或存档编号。"))
    initial = [{"comment_id": r["comment_id"], "评论原文": r.get("original_comment") or r.get("cleaned_comment", "")} for r in comments]
    files.append(_book(output / "01_initial_free_annotation.xlsx", ["comment_id", "评论原文", "使用问题", "自由需求描述", "无法判断", "annotator_id", "annotator_role", "annotated_at", "notes"], initial, instructions="这是独立初始标注。只依据评论原文，用自己的语言填写使用问题和自由需求描述；无法可靠判断时填“是”。本表不提供系统分类答案。"))
    elderly=[]
    for r in comments:
        text=r.get("original_comment") or r.get("cleaned_comment", ""); level, reason=_elderly_candidate(text)
        elderly.append({"comment_id":r["comment_id"],"评论原文":text,"候选提示（非人工结论）":level,"文本依据":reason,"人工判定":"","reviewer_id":"","reviewed_at":"","notes":""})
    files.append(_book(output / "02_elderly_context_candidates.xlsx", list(elderly[0]) if elderly else ["comment_id","评论原文","候选提示（非人工结论）","文本依据","人工判定","reviewer_id","reviewed_at","notes"], elderly, instructions="候选提示仅用于抽样和排期，不是人工结论。人工判定只能依据评论原文填写“明确文本支持 / 一般 / 未知”。"))
    classification=[]
    for c in comments:
        classification.append({"comment_id":c["comment_id"],"评论原文":c.get("original_comment") or c.get("cleaned_comment", ""),"system_requirement_candidates":"","人工分类":"","审核决定":"","reviewer_id":"","reviewed_at":"","review_notes":""})
    files.append(_book(output / "03_later_classification_review.xlsx", list(classification[0]) if classification else [], classification, instructions="仅在独立初始标注封存后开展。system_requirement_candidates 可由研究人员另行填入；审核决定填 approve / modify / reject。不得反向修改初始自由标注。"))
    mapping_headers=[]
    for row in mappings:
        mapping_headers.extend(k for k in row if k not in mapping_headers)
    for key in ("knowledge_source","review_decision","review_status","reviewer_id","reviewer_note","reviewed_at"):
        if key not in mapping_headers: mapping_headers.append(key)
    mapping_rows=[dict(r, review_decision=r.get("review_decision", ""), review_status=r.get("review_status","pending_review")) for r in mappings]
    files.append(_book(output / "04_mapping_review.xlsx", mapping_headers or ["mapping_id","requirement_id","function_id","structure_id","evidence_spans","mapping_reason","derivation_method","knowledge_source","review_decision","review_status","reviewer_id","reviewer_note","reviewed_at"], mapping_rows, instructions="审核需求—功能—结构、原评论证据、两段关系理由和知识来源。review_decision 只填 agree / modify / reject。导入转换：agree→review_status=approved；reject→review_status=rejected；modify→review_status=pending_review，并由研究人员根据 reviewer_note 新建候选，不得改原映射身份和证据。reviewer_id、reviewer_note、reviewed_at 必须由真实审核过程产生。"))
    files.append(_review_book(output / "05_blind_review.xlsx", schemes, "blind"))
    changes=["change_id","v1_scheme_id","v2_scheme_id","reviewer_id","low_dimension","expert_issue","linked_review_comment","evaluated_at","requirement_id","function_id","structure_id","prompt_field","old_value","new_value","reason"]
    files.append(_book(output / "06_v1_v2_changes.xlsx", changes, [], instructions="逐项记录V1到V2的修改、对应低分维度、专家问题、证据链和修改理由。不得填写未实际发生的修改。"))
    files.append(_review_book(output / "07_v1_v2_blind_review.xlsx", schemes, "v1v2"))
    disclosure = "# 规则开发与评价数据边界\n\n本研究的 rules 已查看全部数据，因此该数据上的结果只能描述为开发集/回顾性评价，不能声称独立验证。\n\n若要报告独立验证，必须预先冻结规则，并在未参与规则开发的新数据上评价，同时记录数据隔离、时间边界和人员隔离。\n"
    boundary=output/"BOUNDARY_DISCLOSURE.md"; boundary.write_text(disclosure,encoding="utf-8"); files.append(boundary)
    missing=[]
    if not req_path: missing.append("requirements")
    if not map_path: missing.append("mapping_candidates")
    if not gen_path: missing.append("generation_schemes")
    missing += ["data_source_verification", "independent_initial_annotations", "human_mapping_decisions", "human_scores_and_approvals"]
    source_artifacts = {}
    for stage, path in (("clean", comment_path), ("requirements", req_path), ("mapping", map_path), ("generation", gen_path)):
        if path: source_artifacts[stage] = {"path": str(path.relative_to(run)).replace("\\", "/"), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    exported = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    result={"status":"ready_with_missing_items" if missing else "ready","run_id":manifest.get("run_id",""),"files":[p.name for p in files],"source_artifacts":source_artifacts,"exported_file_sha256":exported,"missing_items":missing,"import_commands":{
        "initial_annotations":f".venv-research\\Scripts\\python.exe experiment\\run_experiment.py --config <config.json> --input <input.csv> --annotations \"{output / '01_initial_free_annotation.xlsx'}\"",
        "mapping_review":f".venv-research\\Scripts\\python.exe experiment\\run_experiment.py --config <config.json> --input <input.csv> --mapping-review \"{output / '04_mapping_review.xlsx'}\"",
    }}
    (output/"materials_manifest.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    return result
