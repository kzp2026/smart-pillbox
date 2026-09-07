from pathlib import Path
from typing import Iterable
import json
import math
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook


DIMENSIONS = [
    "需求匹配度", "适老化友好性", "功能完整性", "操作便利性",
    "结构合理性", "工程可行性", "创新性",
]

REVIEW_HEADERS = {
    "匿名评委 ID": "reviewer_id", "评委背景": "background", "匿名方案 ID": "scheme_id",
    "评价时间": "evaluated_at", "优点": "strengths", "问题": "issues", "建议": "suggestions",
    **{dimension: dimension for dimension in DIMENSIONS},
}


def _safe_scalar(value):
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).replace('"=', '"\'=')
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _write(path: Path, headers: list[str], rows: Iterable[dict]) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append([_safe_scalar(row.get(header, "")) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(path)
    return path


def _blind_rows(schemes: list[dict]) -> list[dict]:
    rows = []
    for scheme in schemes:
        row = {"匿名方案 ID": scheme["scheme_id"], "方案内容": scheme["content"]}
        rows.append(row)
    return rows


def create_templates(directory: Path, comments: list[dict], requirements: list[dict], mappings: list[dict], schemes: list[dict]) -> list[Path]:
    """Create six blank, auditable annotation workbooks.

    Scheme version is deliberately excluded from reviewer workbooks.
    """
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    outputs = []
    outputs.append(_write(directory / "requirement_annotation.xlsx",
                          ["comment_id", "cleaned_comment", "requirement_id", "annotator_id", "annotated_at", "notes"], comments))
    guidelines = [dict(requirement_id=row['requirement_id'], name=row.get('requirement_name', ''),
                       definition=row.get('requirement_description', ''),
                       include='须由具体评论语境支持；候选关键词：'+'、'.join(row.get('keywords', [])),
                       exclude='排除孤立好评、无关语境和被否定的诉求；无需求标记 NO_REQUIREMENT。',
                       example=next((span.get('text', '') for span in row.get('source_evidence_spans', [])), ''),
                       status='候选标注规范，待人工确认') for row in requirements]
    outputs.append(_write(directory / "annotation_guidelines.xlsx",
                          ["requirement_id", "name", "definition", "include", "exclude", "example", "status"], guidelines))
    mapping_headers = list(dict.fromkeys([key for row in mappings for key in row] + ["review_status", "reviewer_id", "reviewed_at", "review_notes"]))
    mapping_rows = [dict(row, review_status=row.get("review_status", "pending_review")) for row in mappings]
    outputs.append(_write(directory / "expert_mapping.xlsx", mapping_headers, mapping_rows))
    blind_headers = ["匿名评委 ID", "评委背景", "匿名方案 ID", "方案内容", *DIMENSIONS, "优点", "问题", "建议", "评价时间"]
    blind_rows = _blind_rows(schemes)
    outputs.append(_write(directory / "abc_blind_review.xlsx", blind_headers, blind_rows))
    outputs.append(_write(directory / "v1_changes.xlsx",
                          ["change_id", "v1_scheme_id", "v2_scheme_id", "reviewer_id", "low_dimension", "expert_issue", "requirement_id", "function_id", "structure_id", "prompt_field", "old_value", "new_value", "reason"], []))
    outputs.append(_write(directory / "v1_v2_blind_review.xlsx", blind_headers, blind_rows))
    return outputs


def _sheet_rows(path: Path):
    sheet = load_workbook(Path(path), data_only=True).active
    headers = [cell.value for cell in sheet[1]]
    if not headers or any(not isinstance(header, str) or not header for header in headers):
        raise ValueError("workbook has invalid headers")
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        row = dict(zip(headers, values))
        if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
            continue
        yield number, row


def read_review_file(path: Path) -> list[dict]:
    """Import a completed Chinese reviewer workbook into the statistics API shape."""
    rows = []
    for number, source in _sheet_rows(path):
        missing_headers = [header for header in REVIEW_HEADERS if header not in source]
        if missing_headers:
            raise ValueError(f"row {number}: missing review columns {missing_headers}")
        row = {target: source[header] for header, target in REVIEW_HEADERS.items()}
        editable = [target for target in REVIEW_HEADERS.values() if target != "scheme_id"]
        if all(row[field] is None or (isinstance(row[field], str) and not row[field].strip()) for field in editable):
            continue
        for dimension in DIMENSIONS:
            value = row[dimension]
            if isinstance(value, float) and math.isfinite(value) and value.is_integer():
                row[dimension] = int(value)
        if any(value is None or (isinstance(value, str) and not value.strip()) for value in row.values()):
            raise ValueError(f"row {number}: partially completed review")
        rows.append(row)
    return rows


def import_annotations(path: Path, comment_ids, requirement_ids) -> list[dict]:
    """Import free-form initial annotations or the legacy coded table.

    The free-form table intentionally has no system requirement label.  Keeping the
    legacy branch allows already collected coded sheets to remain importable.
    """
    sheet = load_workbook(Path(path), data_only=True).active
    headers = [cell.value for cell in sheet[1]]
    free_headers = ["comment_id", "评论原文", "使用问题", "自由需求描述", "无法判断", "annotator_id", "annotator_role", "annotated_at", "notes"]
    if all(header in headers for header in free_headers):
        valid_comments, seen, rows = set(comment_ids), set(), []
        for number, source in _sheet_rows(path):
            editable = ["使用问题", "自由需求描述", "无法判断", "annotator_id", "annotator_role", "annotated_at", "notes"]
            if all(source.get(field) in (None, "") for field in editable):
                continue
            required = ["comment_id", "评论原文", "annotator_id", "annotator_role", "annotated_at"]
            if any(source.get(field) is None or not str(source.get(field)).strip() for field in required):
                raise ValueError(f"row {number}: partially completed free annotation")
            comment_id, annotator_id = str(source["comment_id"]).strip(), str(source["annotator_id"]).strip()
            if comment_id not in valid_comments:
                raise ValueError(f"row {number}: unknown comment_id")
            if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]*", annotator_id):
                raise ValueError(f"row {number}: invalid annotator_id")
            try:
                datetime.fromisoformat(str(source["annotated_at"]).replace("Z", "+00:00"))
            except ValueError as exc:
                raise ValueError(f"row {number}: annotated_at must be ISO 8601") from exc
            key = (comment_id, annotator_id)
            if key in seen:
                raise ValueError(f"row {number}: duplicate annotator and comment")
            seen.add(key)
            undecidable = str(source.get("无法判断") or "").strip()
            if undecidable not in ("是", "否"):
                raise ValueError(f"row {number}: 无法判断 must be 是 or 否")
            issue, free = str(source.get("使用问题") or "").strip(), str(source.get("自由需求描述") or "").strip()
            if undecidable == "否" and not (issue or free):
                raise ValueError(f"row {number}: decidable annotation needs an issue or requirement")
            rows.append({
                "comment_id": comment_id, "original_comment": source["评论原文"], "usage_issue": issue,
                "free_requirement": free, "undecidable": undecidable == "是", "annotator_id": annotator_id,
                "annotator_role": str(source["annotator_role"]).strip(), "annotated_at": source["annotated_at"],
                "notes": source.get("notes") or "",
            })
        return rows

    required = ["comment_id", "cleaned_comment", "requirement_id", "annotator_id", "annotated_at", "notes"]
    valid_comments, valid_requirements = set(comment_ids), set(requirement_ids)
    seen = set()
    rows = []
    for number, source in _sheet_rows(path):
        if any(header not in source for header in required):
            raise ValueError(f"row {number}: invalid annotation columns")
        annotation_fields = ["requirement_id", "annotator_id", "annotated_at", "notes"]
        if all(source[field] is None or (isinstance(source[field], str) and not source[field].strip()) for field in annotation_fields):
            continue
        core = required[:-1]
        if any(source[field] is None or (isinstance(source[field], str) and not source[field].strip()) for field in core):
            raise ValueError(f"row {number}: partially completed annotation")
        comment_id = str(source["comment_id"]).strip()
        annotator_id = str(source["annotator_id"]).strip()
        requirement_id = str(source["requirement_id"]).strip()
        if comment_id not in valid_comments:
            raise ValueError(f"row {number}: unknown comment_id")
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]*", annotator_id):
            raise ValueError(f"row {number}: invalid annotator_id")
        if requirement_id == "NO_REQUIREMENT":
            normalized_requirement = None
        elif requirement_id in valid_requirements:
            normalized_requirement = requirement_id
        else:
            raise ValueError(f"row {number}: unknown requirement_id")
        try:
            datetime.fromisoformat(str(source["annotated_at"]).replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError(f"row {number}: annotated_at must be ISO 8601") from exc
        key = (comment_id, annotator_id)
        if key in seen:
            raise ValueError(f"row {number}: duplicate annotator and comment")
        seen.add(key)
        rows.append({field: source[field] for field in required} | {"comment_id": comment_id, "annotator_id": annotator_id, "requirement_id": normalized_requirement})
    return rows


def import_mapping_reviews(path: Path) -> list[dict]:
    """Normalize the formal decision column to the legacy graph import fields."""
    rows = []
    status_by_decision = {"agree": "approved", "modify": "pending_review", "reject": "rejected"}
    for number, source in _sheet_rows(path):
        decision = str(source.get("review_decision") or "").strip().lower()
        if not decision:
            status=source.get('review_status') or 'pending_review'
            if status not in ('pending_review','approved','rejected'):raise ValueError(f'row {number}: invalid review_status')
            if status in ('approved','rejected'):rows.append(dict(source))
            continue
        if decision not in status_by_decision:
            raise ValueError(f"row {number}: review_decision must be agree, modify, or reject")
        for field in ("mapping_id", "reviewer_id", "reviewer_note", "reviewed_at"):
            if source.get(field) is None or not str(source[field]).strip():
                raise ValueError(f"row {number}: {field} is required")
        try:
            datetime.fromisoformat(str(source["reviewed_at"]).replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError(f"row {number}: reviewed_at must be ISO 8601") from exc
        row = dict(source)
        row["review_decision"] = decision
        row["review_status"] = status_by_decision[decision]
        rows.append(row)
    return rows
