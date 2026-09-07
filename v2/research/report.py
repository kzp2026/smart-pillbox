from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile

from v2.research.analysis import REQUIREMENT_LABELS, SENTIMENT_LABELS
from v2.research.evaluation import DIMENSIONS
from v2.research.provenance import source_snapshot


def csv_bytes(rows: list[dict], columns: list[str] | None = None) -> bytes:
    buffer = io.StringIO(newline='')
    fields = columns or list(rows[0]) if rows else (columns or [])
    writer = csv.DictWriter(buffer, fieldnames=fields, extrasaction='ignore')
    writer.writeheader()
    for row in rows:
        # Spreadsheet formula injection prevention, without altering canonical JSON.
        writer.writerow({k: "'" + v if isinstance(v, str) and v.lstrip().startswith(('=','+','-','@')) else v for k,v in row.items()})
    return buffer.getvalue().encode('utf-8-sig')


def annotation_template(records: list[dict]) -> bytes:
    return csv_bytes([dict(comment_id=r['comment_id'], text=r['text'], task=t, label='', split='test', annotator_id='')
                      for r in records for t in ('sentiment','requirement')])


def review_template(schemes: list[dict]) -> bytes:
    return csv_bytes([dict(scheme_id=s['scheme_id'], reviewer_id='', role='', dimension=d, score='', notes='')
                      for s in schemes for d in DIMENSIONS],
                     ['scheme_id','reviewer_id','role','dimension','score','notes'])


def report_markdown(experiment: dict) -> str:
    card, analysis, manifest = experiment['dataset']['card'], experiment['analysis'], experiment['manifest']
    lines = ['# 论文实验证据报告', '', f"产品：{experiment['product']}", f"运行 ID：{manifest['run_id']}",
             '', '本报告证明记录的分析流程与结果；单凭本报告不能证明算法优于基线、设计有效或论文已达到发表标准。',
             '', '## 数据集与方法', f"原始文件 SHA-256：{card['input_sha256']}", f"清洗数据 SHA-256：{card['dataset_sha256']}",
             f"样本流转：{json.dumps(card['counts'],ensure_ascii=False)}", f"字段映射：{json.dumps(card['column_mapping'],ensure_ascii=False)}",
             f"清洗规则：{json.dumps(card['cleaning'],ensure_ascii=False)}"]
    lines += [f'{k}：{v or "未提供"}' for k,v in card['provenance'].items()]
    lines += ['', '## 实际算法与参数', json.dumps(analysis['methods'],ensure_ascii=False,indent=2),
              '中文字符 2–3 gram TF-IDF 为可复现基线，不是语义嵌入 BERTopic；聚类指标不是分类准确率。',
              f"规则需求证据覆盖率：{analysis['evidence_coverage']:.2%}（命中评论数 / 有效评论数，不代表正确率）。",
              '', '## 独立测试集算法比较']
    if experiment['evaluation']['metrics']:
        for row in experiment['evaluation']['metrics']:
            lines.append(f"{row['task']} / {row['method']}：n={row['n_test']}，Accuracy={row['accuracy']:.4f}，macro P/R/F1={row['macro_precision']:.4f}/{row['macro_recall']:.4f}/{row['macro_f1']:.4f}")
    else: lines.append('未提供真实人工 test 标注，准确率、Precision、Recall、F1 均未计算。')
    lines += experiment['evaluation']['notes']
    lines += ['', '## 真实专家/用户评价']
    if experiment['review_summary']:
        lines += [f"{r['scheme_id']} / {r['role']} / {r['dimension']}：n={r['n']}，均值={r['mean']:.2f}，样本标准差={r['std'] if r['std'] is not None else '不可计算（n<2）'}" for r in experiment['review_summary']]
    else: lines.append('未提供真实评价；系统规则自检不作为专家评分。')
    lines += ['', '## 论文证据检查']
    lines += [f"{r['item']}：{r['status']}。{r['detail']}" for r in experiment['readiness']]
    lines += ['', '## 复现与限制', '本归档含清洗后样本、排除日志、预测、配置、人工输入（如有）和逐文件校验清单。原始上传不重复复制到本归档。',
              '手机号/邮箱自动遮盖不是完整匿名化；公开前人工去除姓名、地址和可识别的罕见描述。',
              '需求类别与功能结构为人工规则，适用性需领域专家确认。不得将一套规则的消融比较宣传为对所有方法的优势。',
              '人工标签/评分的来源由提交者声明，系统只能验证格式和计算；样本量、盲评、多人一致性及统计显著性需研究者论证。',
              '复现：进入解压目录/source，再执行 python -m v2.research.reproduce ../experiment.json；依赖版本见环境清单.txt。',
              f"代码版本：{manifest['version']}；算法源码摘要：{manifest['source_sha256']}",
              '', '## 标注约定', f"sentiment 单选：{' / '.join(SENTIMENT_LABELS)}。", f"requirement 多选用 | 分隔：{' / '.join(REQUIREMENT_LABELS)}。",
              '每评论每任务仅一条仲裁后的真值；annotator_id 填匿名 ID；split 为 train/dev/test。不同方法使用完全相同 test ID。',
              '评分锚点：1=明显不满足，2=较差，3=基本满足，4=较好，5=充分满足。角色为专家或用户；需交代招募、盲评顺序和实验方案。']
    return '\n\n'.join(lines)


def chart_png(experiment: dict) -> bytes:
    # Object-oriented Agg renderer avoids global pyplot state in concurrent sessions.
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure
    figure = Figure(figsize=(8,4.5), dpi=200)
    FigureCanvasAgg(figure)
    axis = figure.add_subplot(111)
    metrics = experiment['evaluation']['metrics']
    if metrics:
        labels = [r['method']+'\n'+r['task'] for r in metrics]
        values = [r['macro_f1'] for r in metrics]
        axis.bar(range(len(values)), values, color='#2166ac')
        axis.set_xticks(range(len(labels)), labels, rotation=20, ha='right', fontsize=8)
        axis.set_ylim(0,1); axis.set_ylabel('Macro F1 (held-out test)')
        axis.set_title('Same test IDs per task; no significance claim')
    else:
        counts = experiment['dataset']['card']['counts']
        axis.bar(list(counts), list(counts.values()), color='#2166ac')
        axis.set_ylabel('Rows'); axis.set_title('Dataset flow (not accuracy or human evaluation)')
    axis.spines[['top','right']].set_visible(False)
    figure.tight_layout()
    buffer = io.BytesIO(); figure.savefig(buffer, format='png', dpi=200)
    return buffer.getvalue()


def report_docx(experiment: dict, chart: bytes) -> bytes:
    # Reuse the deployed Python DOCX engine: no Node subprocess needed on Cloud.
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    style = document.styles['Normal']
    style.font.name = 'Arial'; style.font.size = Pt(10)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    for block in report_markdown(experiment).split('\n\n'):
        if block.startswith('# '): document.add_heading(block[2:], 0)
        elif block.startswith('## '): document.add_heading(block[3:], 1)
        else: document.add_paragraph(block)
    document.add_heading('实验图表',1)
    document.add_picture(io.BytesIO(chart), width=Cm(16))
    document.add_paragraph('图表与 CSV 同源；未有标注时仅展示样本流转。')
    buffer=io.BytesIO(); document.save(buffer)
    return buffer.getvalue()


def build_archive(experiment: dict) -> bytes:
    analysis = experiment['analysis']
    files = {'experiment.json': json.dumps(experiment,ensure_ascii=False,indent=2,allow_nan=False).encode('utf-8'),
             'report.md': report_markdown(experiment).encode('utf-8'),
             'cleaned_comments.csv': csv_bytes(experiment['dataset']['records']),
             'excluded_rows.csv': csv_bytes(experiment['dataset']['exclusions'],['source_row','reason']),
             'predictions.csv': csv_bytes(analysis['predictions']),
             'requirements.csv': csv_bytes(analysis['mappings']), 'topics.csv': csv_bytes(analysis['topics']),
             'keywords.csv': csv_bytes(analysis['keywords']),
             'comparison.csv': csv_bytes(experiment['evaluation']['metrics']),
             'confusion.csv': csv_bytes(experiment['evaluation']['confusion']),
             'gold_labels.csv': csv_bytes(experiment['gold']), 'human_reviews.csv': csv_bytes(experiment['reviews']),
             'review_summary.csv': csv_bytes(experiment['review_summary']),
             'annotation_template.csv': annotation_template(experiment['dataset']['records']),
             'review_template.csv': review_template(experiment['schemes']),
             '环境清单.txt': '\n'.join(f'{k}=={v}' for k,v in experiment['manifest']['dependencies'].items()).encode('utf-8')}
    files['charts.png'] = chart_png(experiment)
    files['论文实验报告.docx'] = report_docx(experiment, files['charts.png'])
    snapshot=source_snapshot()
    if {p:hashlib.sha256(data).hexdigest() for p,data in snapshot.items()} != experiment['manifest']['source_files']:
        raise ValueError('运行期间代码发生变化，请重跑后导出。')
    files.update({'source/'+name:data for name,data in snapshot.items()})
    manifest = dict(experiment['manifest'], files={name:hashlib.sha256(data).hexdigest() for name,data in files.items()})
    files['manifest.json'] = json.dumps(manifest,ensure_ascii=False,indent=2).encode('utf-8')
    buffer=io.BytesIO()
    with zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,data in files.items(): archive.writestr(name,data)
    return buffer.getvalue()
